"""Parse uploaded files into text, keeping page numbers where the format has them.

Dispatch is by file extension. Every parser returns a list of `Page`s; formats
without real pages (DOCX, TXT, MD) return a single page numbered 1.
"""

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import PurePath


@dataclass(frozen=True)
class Page:
    page: int  # 1-based
    text: str


class UnsupportedFileType(ValueError):
    pass


class ParseError(ValueError):
    pass


def _parse_pdf(data: bytes) -> list[Page]:
    import pymupdf

    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            pages = [Page(page=i + 1, text=p.get_text("text")) for i, p in enumerate(doc)]
    except Exception as exc:  # PyMuPDF raises several unrelated types on bad input
        raise ParseError(f"could not read PDF: {exc}") from exc
    return pages


def _parse_docx(data: bytes) -> list[Page]:
    import io

    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"could not read DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    # Tables are common in business docs; keep their cell text, one row per line.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page(page=1, text="\n\n".join(parts))]


def _parse_text(data: bytes) -> list[Page]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return [Page(page=1, text=text)]


PARSERS: dict[str, Callable[[bytes], list[Page]]] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".txt": _parse_text,
    ".md": _parse_text,
}

SUPPORTED_EXTENSIONS = tuple(PARSERS)


def extension_of(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def check_supported(filename: str) -> None:
    ext = extension_of(filename)
    if ext not in PARSERS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or filename}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )


def parse_file(filename: str, data: bytes) -> list[Page]:
    """Parse `data` using the parser for `filename`'s extension.

    Pages with no text are dropped. Raises `UnsupportedFileType` or `ParseError`.
    """
    check_supported(filename)
    pages = PARSERS[extension_of(filename)](data)
    pages = [p for p in pages if p.text.strip()]
    if not pages:
        raise ParseError("no extractable text (scanned PDFs need OCR, which is not supported)")
    return pages
